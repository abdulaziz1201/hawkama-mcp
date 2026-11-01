"""
MCP الحوكمة — يقرأ ملفات مجلد عام في Google Drive (PDF / Word / Google Docs)
ويعطي Claude أدوات للبحث فيها وقراءتها.

- ما يحتاج مفتاح API: يكفي أن يكون المجلد "أي شخص لديه الرابط ← عارض".
- أي ملف تضيفه للمجلد يظهر تلقائياً بعد التحديث الدوري (افتراضياً كل 6 ساعات).
"""
import gzip
import html
import io
import json
import logging
import os
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor

import docx
import pypdfium2 as pdfium
import requests
from mcp.server.fastmcp import FastMCP
from starlette.requests import Request
from starlette.responses import JSONResponse

# ───────────── الإعدادات (تتغيّر من متغيرات البيئة في Render) ─────────────
FOLDER_ID = os.environ.get("FOLDER_ID", "1BXhhu7AqzbrtUnRz48KjZUKBQW4phwId")
SERVER_NAME = os.environ.get("SERVER_NAME", "حوكمة الجمعيات الأهلية")
REFRESH_HOURS = float(os.environ.get("REFRESH_HOURS", "6"))
PORT = int(os.environ.get("PORT", "8000"))

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s")
log = logging.getLogger("hawkama")
http = requests.Session()
http.headers["User-Agent"] = "Mozilla/5.0 (hawkama-mcp)"

mcp = FastMCP(
    SERVER_NAME,
    instructions=(
        "هذا الخادم يحتوي على مكتبة وثائق رسمية واسترشادية عن حوكمة الجمعيات والمؤسسات الأهلية في السعودية "
        "(الأنظمة واللوائح، قواعد الحوكمة، الأدلة الاسترشادية، أدلة التقييم الذاتي حسب حجم الجمعية). "
        "ابدأ بـ search_documents للبحث، ثم read_document لقراءة الصفحات المهمة كاملة. "
        "اذكر دائماً اسم الوثيقة ورقم الصفحة ورابطها عند الإجابة."
    ),
    host="0.0.0.0",
    port=PORT,
    stateless_http=True,
)

# ───────────── تطبيع النص العربي ─────────────
TASHKEEL = re.compile(r"[\u0610-\u061A\u064B-\u065F\u0670\u06D6-\u06ED\u0640]")
NORM_TABLE = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا", "ى": "ي", "ة": "ه"})
STOPWORDS = {"في", "من", "على", "علي", "الى", "الي", "عن", "او", "ما", "هل", "مع", "ان", "التي", "الذي", "هو", "هي", "و"}


def clean(text: str) -> str:
    """نص للعرض: بدون تشكيل أو تطويل، ومسافات موحّدة."""
    text = TASHKEEL.sub("", text)
    return re.sub(r"[ \t\r\f\v]+", " ", text).strip()


def norm(text: str) -> str:
    """نص للبحث: نفس طول clean() حرفاً بحرف، عشان مواضع البحث تطابق نص العرض."""
    return text.translate(NORM_TABLE).lower()


def term_variants(term: str) -> set[str]:
    # مشاكل شائعة في استخراج نص PDF العربي:
    # 1) قلب "لا" إلى "ال" (مثل: الاطلاع ← االطالع)
    # 2) كلمات كاملة مكتوبة بالمقلوب (مثل: مجلس ← سلجم)
    v = {term, term.replace("لا", "ال")}
    return v | {x[::-1] for x in v if len(x) > 2}


# ───────────── قراءة المجلد العام بدون مفتاح API ─────────────
ENTRY_RE = re.compile(r'<a href="([^"]+)"[^>]*>.*?<div class="flip-entry-title">(.*?)</div>', re.S)
ID_RE = re.compile(r"/(?:file|document|spreadsheets|presentation)/d/([\w-]+)|/folders/([\w-]+)|[?&]id=([\w-]+)")


def crawl(folder_id: str, path: str = "", depth: int = 0) -> list[dict]:
    if depth > 6:
        return []
    r = http.get(f"https://drive.google.com/embeddedfolderview?id={folder_id}", timeout=30)
    r.raise_for_status()
    items = []
    for href, title in ENTRY_RE.findall(r.text):
        title = html.unescape(title).strip()
        m = ID_RE.search(href)
        if not m:
            continue
        fid = next(g for g in m.groups() if g)
        if "/folders/" in href:
            items += crawl(fid, f"{path}/{title}" if path else title, depth + 1)
        else:
            kind = next((k for k in ("document", "spreadsheets", "presentation") if f"/{k}/" in href), "file")
            items.append({"id": fid, "title": title, "category": path or "عام", "kind": kind})
    return items


def download(item: dict) -> bytes:
    fid, kind = item["id"], item["kind"]
    url = {
        "document": f"https://docs.google.com/document/d/{fid}/export?format=txt",
        "spreadsheets": f"https://docs.google.com/spreadsheets/d/{fid}/export?format=csv",
        "presentation": f"https://docs.google.com/presentation/d/{fid}/export/txt",
    }.get(kind, f"https://drive.usercontent.google.com/download?id={fid}&export=download&confirm=t")
    r = http.get(url, timeout=120)
    r.raise_for_status()
    return r.content


def extract_pages(data: bytes) -> list[str]:
    """يرجّع قائمة صفحات نصية."""
    if data[:4] == b"%PDF":
        pdf = pdfium.PdfDocument(data)
        try:
            return [clean(pdf[i].get_textpage().get_text_range()) for i in range(len(pdf))]
        finally:
            pdf.close()
    if data[:2] == b"PK":  # Word (.docx)
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8", errors="ignore")
    text = clean(text)
    # نقسم النص الطويل إلى "صفحات" تقريبية بحجم 3000 حرف
    return [text[i : i + 3000] for i in range(0, len(text), 3000)] or [""]


def load_one(item: dict) -> dict:
    doc = dict(item, url=f"https://drive.google.com/file/d/{item['id']}/view", pages=[], error=None)
    try:
        doc["pages"] = extract_pages(download(item))
    except Exception as e:  # ملف تالف أو نوع غير مدعوم
        doc["error"] = str(e)[:200]
        log.warning("تعذّر قراءة %s: %s", item["title"], e)
    doc["norm_pages"] = [norm(p) for p in doc["pages"]]
    text_pages = sum(1 for p in doc["pages"] if len(p) > 40)
    doc["scanned"] = bool(doc["pages"]) and text_pages < len(doc["pages"]) / 2
    return doc


# ───────────── الفهرس في الذاكرة + تحديث دوري ─────────────
# الفهرس يُبنى وقت النشر (Build) ويُحفظ في ملف، عشان التشغيل يكون فوري حتى بعد "نوم" الخدمة المجانية.
# بعد التشغيل، الخادم يضيف الملفات الجديدة فقط (تحديث خفيف) عشان ما يثقل على المعالج.
INDEX_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "index.json.gz")
STATE = {"docs": [], "loaded_at": 0.0, "error": None, "source": None}
READY = threading.Event()
SAVED_KEYS = ("id", "title", "category", "kind", "url", "pages", "error", "scanned")


def save_index(docs: list[dict]):
    with gzip.open(INDEX_FILE, "wt", encoding="utf-8") as f:
        json.dump([{k: d[k] for k in SAVED_KEYS} for d in docs], f, ensure_ascii=False)


def load_index() -> bool:
    if not os.path.exists(INDEX_FILE):
        return False
    with gzip.open(INDEX_FILE, "rt", encoding="utf-8") as f:
        docs = json.load(f)
    for d in docs:
        d["norm_pages"] = [norm(p) for p in d["pages"]]
    STATE.update(docs=docs, loaded_at=os.path.getmtime(INDEX_FILE), source="build")
    READY.set()
    log.info("تم تحميل الفهرس الجاهز: %d ملف", len(docs))
    return True


def build_index(workers: int = 1) -> list[dict]:
    t0 = time.time()
    items = crawl(FOLDER_ID)
    with ThreadPoolExecutor(max_workers=workers) as pool:
        docs = list(pool.map(load_one, items))
    STATE.update(docs=docs, loaded_at=time.time(), error=None, source="live")
    READY.set()
    log.info("تم تحميل %d ملف في %.0f ثانية", len(docs), time.time() - t0)
    return docs


def incremental_refresh():
    """يضيف الملفات الجديدة ويحذف المحذوفة ويحدّث الأسماء، بدون إعادة قراءة كل شيء."""
    items = crawl(FOLDER_ID)
    current = {d["id"]: d for d in STATE["docs"]}
    docs = []
    for it in items:
        d = current.get(it["id"])
        if d is None or d.get("error"):
            d = load_one(it)
            time.sleep(0.5)  # نترك المعالج يتنفس عشان الطلبات ما تتأخر
        else:
            d = dict(d, title=it["title"], category=it["category"])
        docs.append(d)
    added = len([i for i in items if i["id"] not in current])
    STATE.update(docs=docs, loaded_at=time.time(), error=None)
    log.info("تحديث: %d ملف (%d جديد)", len(docs), added)


def refresher():
    if not load_index():
        try:
            build_index(workers=1)
        except Exception as e:
            STATE["error"] = str(e)
            log.exception("فشل التحميل")
    time.sleep(120)  # ننتظر شوي بعد التشغيل قبل أول تحديث
    while True:
        try:
            incremental_refresh()
        except Exception as e:
            STATE["error"] = str(e)
            log.exception("فشل التحديث")
        time.sleep(REFRESH_HOURS * 3600)


def docs_or_wait() -> list[dict] | str:
    if READY.wait(timeout=50):
        return STATE["docs"]
    return "المكتبة ما زالت تُحمَّل (أول تشغيل يأخذ دقيقة أو دقيقتين). أعد المحاولة بعد قليل."


def find_doc(docs: list[dict], document: str) -> dict | None:
    for d in docs:
        if d["id"] == document:
            return d
    q = norm(clean(document))
    matches = [d for d in docs if q in norm(clean(d["title"]))]
    return min(matches, key=lambda d: len(d["title"])) if matches else None


# ───────────── الأدوات ─────────────
@mcp.tool()
def list_documents(category: str = "") -> list[dict] | str:
    """
    يعرض كل الوثائق المتاحة في المكتبة مع تصنيفها (مثل: الأدلة العامة، الجمعيات الصغيرة، الجمعيات الكبيرة...).
    category: اختياري، لتصفية الوثائق حسب التصنيف أو جزء منه.
    """
    docs = docs_or_wait()
    if isinstance(docs, str):
        return docs
    out = []
    for d in sorted(docs, key=lambda d: (d["category"], d["title"])):
        if category and norm(category) not in norm(d["category"]):
            continue
        out.append({
            "id": d["id"], "العنوان": d["title"], "التصنيف": d["category"],
            "عدد_الصفحات": len(d["pages"]), "الرابط": d["url"],
            **({"ملاحظة": "ملف مصوّر، النص غير قابل للبحث"} if d["scanned"] else {}),
            **({"خطأ": d["error"]} if d["error"] else {}),
        })
    return out


@mcp.tool()
def search_documents(query: str, category: str = "", max_results: int = 8) -> list[dict] | str:
    """
    يبحث داخل نصوص كل الوثائق ويرجّع أكثر المقاطع صلة مع اسم الوثيقة ورقم الصفحة والرابط.
    query: كلمات البحث، مثل "تعارض المصالح" أو "نصاب اجتماع الجمعية العمومية".
    category: اختياري، لحصر البحث في تصنيف معيّن مثل "الجمعيات الصغيرة".
    """
    docs = docs_or_wait()
    if isinstance(docs, str):
        return docs
    terms = [t for t in norm(clean(query)).split() if len(t) > 1 and t not in STOPWORDS]
    if not terms:
        return "اكتب كلمة بحث أوضح."
    variants = [term_variants(t) for t in terms]
    hits = []
    for d in docs:
        if category and norm(category) not in norm(d["category"]):
            continue
        title_n = norm(clean(d["title"]))
        for i, page in enumerate(d["norm_pages"]):
            per_term = []  # (عدد التكرار، أول موضع) لكل كلمة مطابقة
            for vs in variants:
                c, pos = 0, None
                for v in vs:
                    n = page.count(v)
                    if n:
                        c += n
                        p = page.find(v)
                        pos = p if pos is None else min(pos, p)
                if c:
                    per_term.append((c, pos))
            if per_term:
                # المقتطف يبدأ عند أندر كلمة في الصفحة (غالباً هي الأهم مثل "نصاب")
                anchor = min(per_term)[1]
                title_bonus = sum(1 for vs in variants if any(v in title_n for v in vs))
                total = sum(c for c, _ in per_term)
                hits.append((len(per_term), title_bonus, total, d, i, anchor))
    hits.sort(key=lambda h: (h[0], h[1], h[2]), reverse=True)
    results = []
    for found, _, _, d, i, pos in hits[: max(1, min(max_results, 20))]:
        text = d["pages"][i]
        start = max(0, pos - 250)
        results.append({
            "الوثيقة": d["title"], "id": d["id"], "التصنيف": d["category"], "الصفحة": i + 1,
            "الكلمات_المطابقة": f"{found}/{len(terms)}",
            "مقتطف": ("…" if start else "") + text[start : start + 700] + ("…" if start + 700 < len(text) else ""),
            "الرابط": d["url"],
        })
    return results or "ما لقيت نتائج. جرّب كلمات أخرى أو مرادفات."


@mcp.tool()
def read_document(document: str, from_page: int = 1, to_page: int = 0) -> dict | str:
    """
    يقرأ نص وثيقة كاملة أو صفحات محددة منها.
    document: معرّف الوثيقة (id) أو جزء من عنوانها.
    from_page / to_page: نطاق الصفحات (اتركها فارغة لقراءة الوثيقة من البداية).
    """
    docs = docs_or_wait()
    if isinstance(docs, str):
        return docs
    d = find_doc(docs, document)
    if not d:
        return "ما لقيت وثيقة بهذا الاسم. استخدم list_documents لعرض الأسماء."
    n = len(d["pages"])
    a = max(1, from_page)
    b = min(n, to_page if to_page >= a else n)
    parts, size, last = [], 0, a - 1
    for p in range(a, b + 1):
        chunk = f"\n--- صفحة {p} ---\n{d['pages'][p - 1]}"
        if size + len(chunk) > 40000 and parts:
            break
        parts.append(chunk)
        size += len(chunk)
        last = p
    return {
        "الوثيقة": d["title"], "التصنيف": d["category"], "الرابط": d["url"],
        "عدد_الصفحات": n, "الصفحات_المعروضة": f"{a}-{last}",
        **({"للمتابعة": f"اقرأ من الصفحة {last + 1}"} if last < b else {}),
        "النص": "".join(parts),
    }


@mcp.custom_route("/", methods=["GET"])
async def health(_: Request) -> JSONResponse:
    return JSONResponse({
        "name": SERVER_NAME, "ready": READY.is_set(), "documents": len(STATE["docs"]),
        "loaded_at": time.strftime("%Y-%m-%d %H:%M", time.gmtime(STATE["loaded_at"])) if STATE["loaded_at"] else None,
        "source": STATE["source"], "mcp_endpoint": "/mcp", "error": STATE["error"],
    })


if __name__ == "__main__":
    import sys
    if "--build" in sys.argv:
        # يُستدعى وقت النشر: يبني الفهرس ويحفظه في ملف ثم يخرج
        save_index(build_index(workers=6))
    else:
        threading.Thread(target=refresher, daemon=True).start()
        mcp.run(transport="streamable-http")
